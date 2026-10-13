"""Exercise the real installer and model-selected shell against an unrelated environment."""

from __future__ import annotations

import json
import os
import shlex
import subprocess
import sys
import tempfile
import threading
import zipfile
from contextlib import contextmanager
from http.server import BaseHTTPRequestHandler, ThreadingHTTPServer
from pathlib import Path

ROOT = Path(__file__).resolve().parents[1]


def main() -> None:
    with tempfile.TemporaryDirectory(prefix="feather-install-check-") as temporary:
        root = Path(temporary)
        caller = root / "caller project"
        caller.mkdir()
        (caller / "pyproject.toml").write_text('[project]\nname="unrelated"\nversion="0.0.0"\n')
        subprocess.run([sys.executable, "-m", "venv", "--without-pip", str(caller / ".venv")], check=True)
        poison = caller / "foreign packages"
        poison.mkdir()
        (poison / "docx.py").write_text('raise AssertionError("Caller Python leaked into Feather")\n')
        before = set(caller.rglob("*"))
        install_root = root / "Feather environment"
        binary_dir = root / "bin space"
        environment = {
            **os.environ,
            "FEATHER_INSTALL_DIR": str(install_root),
            "FEATHER_BIN_DIR": str(binary_dir),
            "FEATHER_SOURCE_DIR": str(ROOT),
            "VIRTUAL_ENV": str(caller / ".venv"),
            "UV_PROJECT_ENVIRONMENT": str(caller / ".venv"),
            "UV_PROJECT": str(caller),
            "PYTHONPATH": str(poison),
            "PYTHONHOME": str(poison),
            "PATH": str(caller / ".venv/bin") + os.pathsep + os.environ["PATH"],
        }
        install(environment, caller)
        launcher = binary_dir / "feather"
        subprocess.run([str(launcher), "--help"], env=environment, cwd=caller, check=True, capture_output=True)
        first = (install_root / "current").resolve()

        # Remote mode is exercised against a real Git repository, without requiring GitHub access.
        remote_environment = {**environment, "FEATHER_REPO": str(ROOT), "FEATHER_REF": "HEAD"}
        remote_environment.pop("FEATHER_SOURCE_DIR")
        install(remote_environment, caller)
        assert (install_root / "current").resolve() != first
        assert first.is_dir(), "Successful updates should retain the old environment"
        selected = (install_root / "current").resolve()
        failed = install({**environment, "FEATHER_SOURCE_DIR": str(root / "missing")}, caller, check=False)
        assert failed.returncode != 0
        assert (install_root / "current").resolve() == selected, "Failed updates must preserve the installation"
        assert not (install_root / ".install-lock").exists()
        assert set(caller.rglob("*")) == before, "Installation wrote into the caller's project"

        workspace = root / "task workspace"
        workspace.mkdir()
        program = """
import json, sys
from importlib.metadata import version
from pathlib import Path
import docx, pptx, openpyxl
result = {
    'prefix': sys.prefix,
    'modules': [docx.__file__, pptx.__file__, openpyxl.__file__],
    'versions': {name: version(name) for name in ('paper-docx', 'paper-pptx', 'paper-xlsx')},
}
Path(sys.argv[1]).write_text(json.dumps(result))
docx.Document().save('brief.docx')
pptx.Presentation().save('deck.pptx')
openpyxl.Workbook().save('sheet.xlsx')
print('Paper environment verified')
"""
        command = " && ".join(
            f"{executable} -c {shlex.quote(program)} {filename}"
            for executable, filename in (("python", "python.json"), ("python3", "python3.json"))
        )
        with model_server(command) as (endpoint, requests, started, release):
            process = subprocess.Popen(
                [str(launcher), "--model", "test", "--base-url", endpoint,
                 "--workdir", str(workspace), "Verify the Paper Python environment."],
                cwd=caller,
                env={**environment, "OPENAI_API_KEY": "synthetic-test-key"},
                stdout=subprocess.PIPE, stderr=subprocess.PIPE, text=True,
            )
            try:
                assert started.wait(30), "CLI did not reach the model endpoint"
                install(remote_environment, caller)
                assert (install_root / "current").resolve() != selected
            finally:
                release.set()
            stdout, stderr = process.communicate(timeout=60)
        assert process.returncode == 0, stderr
        result = json.loads(stdout)
        assert result["final_answer"] == "Verified."
        assert len(requests) == 2
        assert "Paper environment verified" in json.dumps(requests[1]["messages"])
        assert ".agents/docx/SKILL.md" in requests[0]["messages"][0]["content"]
        for name in ("python.json", "python3.json"):
            evidence = json.loads((workspace / name).read_text())
            assert Path(evidence["prefix"]).resolve() == selected / "venv"
            for module in evidence["modules"]:
                assert Path(module).resolve().is_relative_to(selected / "venv")
            assert evidence["versions"] == {"paper-docx": "0.2.0", "paper-pptx": "0.2.0", "paper-xlsx": "0.2.1"}
        for name, member in (("brief.docx", "word/document.xml"), ("deck.pptx", "ppt/presentation.xml"), ("sheet.xlsx", "xl/workbook.xml")):
            with zipfile.ZipFile(workspace / name) as archive:
                assert member in archive.namelist()
        assert set(caller.rglob("*")) == before, "Running Feather changed the unrelated project"
        print("Installer verified: isolated packages, local/remote sources, updates, failure recovery, real Python/Paper tool execution.")


def install(environment: dict[str, str], caller: Path, *, check: bool = True) -> subprocess.CompletedProcess[str]:
    result = subprocess.run(
        ["bash", str(ROOT / "install.sh")], env=environment, cwd=caller,
        capture_output=True, text=True, timeout=300, check=False,
    )
    if check and result.returncode:
        raise AssertionError(result.stdout + result.stderr)
    return result


@contextmanager
def model_server(command: str):
    requests = []
    started = threading.Event()
    release = threading.Event()

    class Handler(BaseHTTPRequestHandler):
        def do_POST(self):
            request = json.loads(self.rfile.read(int(self.headers["Content-Length"])))
            requests.append(request)
            delta = {"role": "assistant", "content": "Verified."}
            finish = "stop"
            if len(requests) == 1:
                started.set()
                if not release.wait(60):
                    raise TimeoutError("Update check did not release the model response")
                delta = {"role": "assistant", "tool_calls": [{
                    "index": 0, "id": "environment-check", "type": "function",
                    "function": {"name": "bash", "arguments": json.dumps({"command": command})},
                }]}
                finish = "tool_calls"
            payload = {
                "id": "check", "object": "chat.completion.chunk", "created": 0,
                "model": "test", "choices": [{"index": 0, "delta": delta, "finish_reason": finish}],
            }
            body = (f"data: {json.dumps(payload)}\n\ndata: [DONE]\n\n").encode()
            self.send_response(200)
            self.send_header("Content-Type", "text/event-stream")
            self.send_header("Content-Length", str(len(body)))
            self.end_headers()
            self.wfile.write(body)

        def log_message(self, *args):
            pass

    server = ThreadingHTTPServer(("127.0.0.1", 0), Handler)
    thread = threading.Thread(target=server.serve_forever, daemon=True)
    thread.start()
    try:
        yield f"http://127.0.0.1:{server.server_port}/v1", requests, started, release
    finally:
        release.set()
        server.shutdown()
        server.server_close()
        thread.join()


if __name__ == "__main__":
    main()
